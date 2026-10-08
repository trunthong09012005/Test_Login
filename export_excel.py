import xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

TEST_MAPPING = {
    "test_tc01_empty_username": {"id": "TC01", "type": "Negative", "desc": "Để trống tên đăng nhập", "steps": "1. Mở trang login\n2. Để trống ô Username\n3. Nhập Password\n4. Click nút Login", "data": "User: (Trống)\nPass: [MASKED]", "expected": "Hiển thị cảnh báo: 'Bạn chưa nhập tên đăng nhập'"},
    "test_tc02_empty_password": {"id": "TC02", "type": "Negative", "desc": "Để trống mật khẩu", "steps": "1. Mở trang login\n2. Nhập Username\n3. Để trống Password\n4. Click nút Login", "data": "User: [TEST_USER]\nPass: (Trống)", "expected": "Hiển thị cảnh báo: 'Bạn chưa nhập mật khẩu'"},
    "test_tc03_valid_username_wrong_password": {"id": "TC03", "type": "Negative", "desc": "Đúng Username, sai Password", "steps": "1. Nhập đúng Username\n2. Nhập sai Password\n3. Click Login", "data": "User: Hợp lệ\nPass: Sai", "expected": "Thông báo lỗi mật khẩu không đúng"},
    "test_tc04_wrong_username_valid_password": {"id": "TC04", "type": "Negative", "desc": "Sai Username, đúng Password", "steps": "1. Nhập sai Username\n2. Nhập đúng Password\n3. Click Login", "data": "User: Sai\nPass: Hợp lệ", "expected": "Thông báo lỗi tài khoản không đúng"},
    "test_tc05_remember_me_persists_after_browser_restart": {"id": "TC05", "type": "Positive", "desc": "Đăng nhập chọn 'Giữ tôi luôn đăng nhập'", "steps": "1. Nhập User/Pass hợp lệ\n2. Tích 'Giữ tôi luôn đăng nhập'\n3. Đóng và mở lại trình duyệt", "data": "User/Pass: Hợp lệ", "expected": "Trình duyệt lưu phiên, không cần đăng nhập lại"},
    "test_tc06_without_remember_me_requires_login_after_restart": {"id": "TC06", "type": "Positive", "desc": "Đăng nhập không chọn giữ đăng nhập", "steps": "1. Nhập User/Pass\n2. Không tích 'Giữ đăng nhập'\n3. Đóng và mở lại", "data": "User/Pass: Hợp lệ", "expected": "Phiên kết thúc, yêu cầu đăng nhập lại"},
    "test_tc07_empty_username_and_password": {"id": "TC07", "type": "Negative", "desc": "Để trống cả Username và Password", "steps": "1. Mở trang login\n2. Để trống cả 2 ô\n3. Click Login", "data": "User/Pass: (Trống)", "expected": "Hiển thị cảnh báo yêu cầu nhập thông tin"},
    "test_tc08_wrong_username_and_password": {"id": "TC08", "type": "Negative", "desc": "Sai cả Username và Password", "steps": "1. Nhập sai Username\n2. Nhập sai Password\n3. Click Login", "data": "User/Pass: Sai", "expected": "Thông báo lỗi tài khoản/mật khẩu không đúng"},
    "test_tc09_leading_whitespace_in_username": {"id": "TC09", "type": "Negative", "desc": "Username có khoảng trắng ở đầu", "steps": "1. Nhập Username có dấu cách ở đầu\n2. Click Login", "data": "User: ' admin'\nPass: Hợp lệ", "expected": "Xử lý khoảng trắng hoặc báo lỗi hợp lý"},
    "test_tc10_trailing_whitespace_in_username": {"id": "TC10", "type": "Negative", "desc": "Username có khoảng trắng ở cuối", "steps": "1. Nhập Username có dấu cách ở cuối\n2. Click Login", "data": "User: 'admin '\nPass: Hợp lệ", "expected": "Xử lý khoảng trắng hoặc báo lỗi hợp lý"},
    "test_tc11_leading_whitespace_in_password": {"id": "TC11", "type": "Negative", "desc": "Password có khoảng trắng ở đầu", "steps": "1. Nhập Password có dấu cách ở đầu", "data": "Pass: ' 123456'", "expected": "Thông báo sai mật khẩu"},
    "test_tc12_trailing_whitespace_in_password": {"id": "TC12", "type": "Negative", "desc": "Password có khoảng trắng ở cuối", "steps": "1. Nhập Password có dấu cách ở cuối", "data": "Pass: '123456 '", "expected": "Thông báo sai mật khẩu"},
    "test_tc13_uppercase_username": {"id": "TC13", "type": "Negative", "desc": "Username viết hoa toàn bộ", "steps": "1. Nhập Username in hoa\n2. Click Login", "data": "User: ADMIN", "expected": "Hệ thống tự động xử lý chữ hoa thành chữ thường hoặc báo lỗi"},
    "test_tc14_wrong_password_case": {"id": "TC14", "type": "Negative", "desc": "Password sai chữ hoa/thường", "steps": "1. Nhập Password sai CapsLock\n2. Click Login", "data": "Pass: sai HOA thường", "expected": "Thông báo sai mật khẩu"},
    "test_tc15_special_characters_in_username": {"id": "TC15", "type": "Negative", "desc": "Username chứa ký tự đặc biệt", "steps": "1. Nhập Username chứa ký tự đặc biệt", "data": "User: @#$%", "expected": "Từ chối đăng nhập"},
    "test_tc16_special_characters_in_password": {"id": "TC16", "type": "Negative", "desc": "Password chứa ký tự đặc biệt", "steps": "1. Nhập Password chứa ký tự đặc biệt", "data": "Pass: @#$%", "expected": "Từ chối đăng nhập"},
    "test_tc17_username_at_excessive_length": {"id": "TC17", "type": "Negative", "desc": "Username dài 256 ký tự", "steps": "1. Nhập Username 256 ký tự", "data": "User: 256 ký tự", "expected": "Hệ thống cắt bớt hoặc báo lỗi"},
    "test_tc18_password_at_excessive_length": {"id": "TC18", "type": "Negative", "desc": "Password dài 256 ký tự", "steps": "1. Nhập Password 256 ký tự", "data": "Pass: 256 ký tự", "expected": "Hệ thống cắt bớt hoặc báo lỗi"},
    "test_tc19_password_is_masked": {"id": "TC19", "type": "Positive", "desc": "Mật khẩu được che khi nhập", "steps": "1. Gõ ký tự vào ô Password", "data": "Pass: 123456", "expected": "Hiển thị dưới dạng dấu sao/chấm đen"},
    "test_tc20_submit_login_with_enter": {"id": "TC20", "type": "Positive", "desc": "Đăng nhập bằng phím Enter", "steps": "1. Nhập User/Pass\n2. Bấm phím Enter", "data": "User/Pass: Hợp lệ", "expected": "Đăng nhập thành công"},
    "test_tc21_five_consecutive_failed_logins": {"id": "TC21", "type": "Negative", "desc": "Đăng nhập sai 5 lần liên tiếp", "steps": "1. Đăng nhập sai 5 lần", "data": "User: Hợp lệ, Pass: Sai", "expected": "Khóa tài khoản tạm thời hoặc hiển thị Captcha"},
    "test_tc22_logout_after_successful_login": {"id": "TC22", "type": "Positive", "desc": "Đăng xuất sau khi đăng nhập", "steps": "1. Đăng nhập\n2. Bấm nút Đăng xuất", "data": "User/Pass: Hợp lệ", "expected": "Đăng xuất thành công, về trang Login"},
    "test_tc23_protected_page_redirects_anonymous_user": {"id": "TC23", "type": "Negative", "desc": "Truy cập trang bảo vệ khi chưa đăng nhập", "steps": "1. Truy cập thẳng URL Dashboard", "data": "Không đăng nhập", "expected": "Chuyển hướng về trang Login"},
    "test_tc24_sql_injection_is_rejected": {"id": "TC24", "type": "Negative", "desc": "SQL Injection vào form", "steps": "1. Nhập mã SQL Injection", "data": "User: ' OR '1'='1", "expected": "Từ chối đăng nhập, không lỗi server"},
    "test_tc25_script_injection_is_not_executed": {"id": "TC25", "type": "Negative", "desc": "XSS Injection vào form", "steps": "1. Nhập mã JS", "data": "User: <script>alert(1)</script>", "expected": "Không thực thi script"},
    "test_tc26_back_after_logout_does_not_restore_session": {"id": "TC26", "type": "Positive", "desc": "Nút Back không khôi phục phiên", "steps": "1. Đăng xuất\n2. Ấn Back trên trình duyệt", "data": "User/Pass: Hợp lệ", "expected": "Vẫn ở trang Login, không cho truy cập"}
}

def export_to_excel():
    tree = ET.parse('reports/junit.xml')
    root = tree.getroot()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Results"
    
    # Header Setup
    headers = ['STT', 'Mã TC', 'Phân loại', 'Mô tả', 'Các bước', 'Dữ liệu kiểm thử', 'Kết quả mong đợi', 'Kết quả thực tế', 'Trạng thái', 'Ghi chú']
    ws.append(headers)
    
    # Styling variables
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    wrap_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    for col in range(1, 11):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        
    testsuite = root.find('testsuite')
    if testsuite is None:
        testsuite = root
        
    count = 1
    for testcase in testsuite.iter('testcase'):
        name = testcase.get('name')
        
        failure = testcase.find('failure')
        if failure is not None:
            status = 'FAILED'
            actual_res = failure.get('message', '').split('\n')[0][:100] + "..."
            note = "Lỗi ghi nhận từ pytest"
        else:
            status = 'PASSED'
            actual_res = "Hệ thống hoạt động đúng như mong đợi"
            note = ""
            
        tc_info = TEST_MAPPING.get(name, {})
        tc_id = tc_info.get("id", f"TC{count:02d}")
        tc_type = tc_info.get("type", "Unknown")
        tc_desc = tc_info.get("desc", name)
        tc_steps = tc_info.get("steps", "")
        tc_data = tc_info.get("data", "")
        tc_expected = tc_info.get("expected", "")
        
        row_data = [count, tc_id, tc_type, tc_desc, tc_steps, tc_data, tc_expected, actual_res, status, note]
        ws.append(row_data)
        
        row_idx = count + 1
        for col_idx in range(1, 11):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if col_idx in [1, 2, 3, 9]:
                cell.alignment = center_align
            else:
                cell.alignment = wrap_align
                
        # Status coloring
        status_cell = ws.cell(row=row_idx, column=9)
        if status == 'FAILED':
            status_cell.fill = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid")
            status_cell.font = Font(color="9C0006", bold=True)
        else:
            status_cell.fill = PatternFill(start_color="EBF1DE", end_color="EBF1DE", fill_type="solid")
            status_cell.font = Font(color="006100", bold=True)
            
        count += 1
        
    # Set column widths
    widths = {'A': 5, 'B': 8, 'C': 12, 'D': 30, 'E': 35, 'F': 20, 'G': 30, 'H': 30, 'I': 15, 'J': 25}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
        
    wb.save('reports/Test_Results_v2.xlsx')
    print("Exported successfully to reports/Test_Results_v2.xlsx")

if __name__ == "__main__":
    export_to_excel()
